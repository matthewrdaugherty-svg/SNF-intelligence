"""Export the DuckDB tables to a filterable/pivotable Excel workbook with lineage and limitations.
v2 — incorporates fixes from docs/VALIDATION_REPORT_01.md.
Usage: python3 scripts/03_export_xlsx.py
"""
import duckdb, pandas as pd, json, datetime, re, csv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
ROOT="/home/claude/snf-intel"; OUT=f"{ROOT}/output/SNF_Intelligence_CareCompare_Sep2026.xlsx"
con=duckdb.connect(f"{ROOT}/db/snf.duckdb", read_only=True)
df=lambda s: con.execute(s).df()
FN = lambda label: f"{label} — footnote"   # consistent, unique footnote header pattern

# ---------- curated master: (db column, friendly header) ----------
CUR = [
 ("ccn","CCN"),("provider_name","Facility name"),("provider_address","Address"),("city_town","City"),("state","State"),("zip_code","ZIP"),
 ("county_parish","County"),("urban","Urban (Y/N)"),("latitude","Latitude"),("longitude","Longitude"),("telephone_number","Phone"),
 ("provider_type","Provider type (Medicare/Medicaid)"),("provider_resides_in_hospital","Located in a hospital (Y/N)"),
 ("ownership_type","Ownership type"),("chain_name","Chain name"),("number_of_facilities_in_chain","Facilities in chain"),
 ("chain_average_overall_5_star_rating","Chain avg overall stars"),("legal_business_name","Legal business name"),
 ("own_data_available","Ownership data available (Y/N)"),
 ("own_direct_owners","Direct owners (orgs & individuals)"),("own_indirect_owners","Indirect owners"),("own_partnership_owners","Partnership-interest owners"),
 ("own_operational_control_orgs","Operational/managerial control — organizations"),("own_operational_control_individuals","Operational/managerial control — individuals"),
 ("own_most_recent_association_date","Most recent owner/manager association date"),
 ("provider_changed_ownership_in_last_12_months","Changed ownership last 12 mo (Y/N)"),
 ("number_of_certified_beds","Certified beds"),("average_number_of_residents_per_day","Avg residents per day"),("average_number_of_residents_per_day_footnote",FN("Avg residents per day")),
 ("date_first_approved_to_provide_medicare_and_medicaid_services","First certified date"),("continuing_care_retirement_community","CCRC (Y/N)"),
 ("special_focus_status","Special Focus status"),("abuse_icon","Abuse icon (Y/N)"),("high_performing_icon","High performing icon (Y/N)"),
 ("most_recent_health_inspection_more_than_2_years_ago","Last health inspection >2 yrs ago (Y/N)"),
 ("overall_rating","Overall star rating"),("overall_rating_footnote",FN("Overall star rating")),
 ("health_inspection_rating","Health inspection stars"),("health_inspection_rating_footnote",FN("Health inspection stars")),
 ("qm_rating","Quality measure stars"),("qm_rating_footnote",FN("Quality measure stars")),
 ("long_stay_qm_rating","Long-stay QM stars"),("long_stay_qm_rating_footnote",FN("Long-stay QM stars")),
 ("short_stay_qm_rating","Short-stay QM stars"),("short_stay_qm_rating_footnote",FN("Short-stay QM stars")),
 ("staffing_rating","Staffing stars"),("staffing_rating_footnote",FN("Staffing stars")),
 ("reported_total_nurse_staffing_hours_per_resident_per_day","Total nurse hrs per resident-day (reported)"),
 ("reported_rn_staffing_hours_per_resident_per_day","RN hrs per resident-day (reported)"),
 ("reported_staffing_footnote",FN("Reported staffing hours")),
 ("adjusted_total_nurse_staffing_hours_per_resident_per_day","Total nurse hrs per resident-day (case-mix adjusted)"),
 ("total_number_of_nurse_staff_hours_per_resident_per_day_on_the_weekend","Weekend total nurse hrs per resident-day"),
 ("reported_physical_therapist_staffing_hours_per_resident_per_day","PT hrs per resident-day"),("physical_therapist_staffing_footnote",FN("PT hrs")),
 ("total_nursing_staff_turnover","Total nursing staff turnover (%)"),("total_nursing_staff_turnover_footnote",FN("Total nursing turnover")),
 ("registered_nurse_turnover","RN turnover (%)"),("registered_nurse_turnover_footnote",FN("RN turnover")),
 ("number_of_administrators_who_have_left_the_nursing_home","Administrators who left"),("administrator_turnover_footnote",FN("Administrators who left")),
 ("nursing_case_mix_index","Nursing case-mix index"),
 ("total_weighted_health_survey_score","Weighted health survey score (lower = better)"),
 ("ss_c1_health_survey_date","Most recent standard health survey date"),("ss_c1_health_deficiencies","Health deficiencies, most recent cycle"),
 ("ss_c1_fire_deficiencies","Fire safety deficiencies, most recent cycle"),
 ("hc_n_total","Health deficiencies, all 3 cycles"),("hc_n_severity_G_or_worse","Deficiencies severity G–L (actual harm or immediate jeopardy), 3 cycles"),
 ("hc_n_actual_harm_GHI","Actual-harm deficiencies (G/H/I), 3 cycles"),("hc_n_immediate_jeopardy_JKL","Immediate-jeopardy deficiencies (J/K/L), 3 cycles"),
 ("hc_n_abuse_neglect","Abuse/neglect/exploitation deficiencies, 3 cycles"),
 ("hc_n_standard","Standard-survey deficiencies, 3 cycles"),("hc_n_complaint","Complaint-survey deficiencies, 3 cycles"),
 ("hc_n_complaint_last_12m","Complaint-survey deficiencies, last 12 months"),("hc_n_complaint_last_12m_G_or_worse","Complaint deficiencies severity G–L, last 12 months"),
 ("sd_n_health_complaint_surveys_3cycles","Health complaint inspections (with citations), 3 cycles"),("sd_n_health_complaint_surveys_last_12m","Health complaint inspections (with citations), last 12 months"),
 ("hc_n_infection_control","Infection-control inspection deficiencies"),("sd_n_infection_control_surveys","Infection-control inspections (with citations)"),
 ("hc_n_last_12m","Health deficiencies, last 12 months"),("hc_n_last_12m_G_or_worse","Deficiencies severity G–L, last 12 months"),
 ("fc_n_total","Fire safety deficiencies, 3 cycles"),
 ("pen_n_fines","Number of fines (3-yr window)"),("pen_total_fines_usd","Total fines $ (3-yr window)"),("pen_n_fines_last_12m","Fines, last 12 months"),
 ("pen_fines_usd_last_12m","Fines $, last 12 months"),("pen_n_payment_denials","Payment denials (3-yr)"),("pen_total_denial_days","Payment denial days (3-yr)"),
 ("pen_last_penalty_date","Most recent penalty date"),
 ("qm_521_adj","QM 521 Short-stay rehospitalization (%, adjusted)"),("qm_521_fn",FN("QM 521")),
 ("qm_522_adj","QM 522 Short-stay outpatient ED visit (%, adjusted)"),("qm_522_fn",FN("QM 522")),
 ("qm_551_adj","QM 551 Long-stay hospitalizations per 1,000 resident-days"),("qm_551_fn",FN("QM 551")),
 ("qm_552_adj","QM 552 Long-stay ED visits per 1,000 resident-days"),("qm_552_fn",FN("QM 552")),
 ("qm_434_4q","QM 434 Short-stay new antipsychotic (%)"),("qm_434_fn",FN("QM 434")),
 ("qm_430_4q","QM 430 Short-stay pneumococcal vaccine (%)"),("qm_430_fn",FN("QM 430")),
 ("qm_472_4q","QM 472 Short-stay flu vaccine (%)"),("qm_472_fn",FN("QM 472")),
 ("qm_401_4q","QM 401 Long-stay ADL help increased (%)"),("qm_401_fn",FN("QM 401")),
 ("qm_410_4q","QM 410 Long-stay falls with major injury (%)"),("qm_410_fn",FN("QM 410")),
 ("qm_479_4q","QM 479 Long-stay pressure ulcers (%)"),("qm_479_fn",FN("QM 479")),
 ("qm_481_4q","QM 481 Long-stay antipsychotic (%)"),("qm_481_fn",FN("QM 481")),
 ("qm_407_4q","QM 407 Long-stay UTI (%)"),("qm_407_fn",FN("QM 407")),
 ("qm_406_4q","QM 406 Long-stay catheter (%)"),("qm_406_fn",FN("QM 406")),
 ("qm_451_4q","QM 451 Long-stay walking worsened (%)"),("qm_451_fn",FN("QM 451")),
 ("qm_404_4q","QM 404 Long-stay excess weight loss (%)"),("qm_404_fn",FN("QM 404")),
 ("qm_408_4q","QM 408 Long-stay depressive symptoms (%)"),("qm_408_fn",FN("QM 408")),
 ("qm_409_4q","QM 409 Long-stay physically restrained (%)"),("qm_409_fn",FN("QM 409")),
 ("qm_415_4q","QM 415 Long-stay pneumococcal vaccine (%)"),("qm_415_fn",FN("QM 415")),
 ("qm_454_4q","QM 454 Long-stay flu vaccine (%)"),("qm_454_fn",FN("QM 454")),
 ("qm_452_4q","QM 452 Long-stay antianxiety/hypnotic (%)"),("qm_452_fn",FN("QM 452")),
 ("qm_480_4q","QM 480 Long-stay new/worsened incontinence (%)"),("qm_480_fn",FN("QM 480")),
 ("qrp_004_ppr_pd_rsrr","QRP 004 Potentially preventable 30-day post-discharge readmission, risk-standardized (%)"),("qrp_004_ppr_pd_rsrr_fn",FN("QRP 004 rate")),
 ("qrp_004_ppr_pd_comp_perf","QRP 004 readmission vs national (CMS determination)"),("qrp_004_ppr_pd_volume","QRP 004 eligible stays"),
 ("qrp_005_dtc_rs_rate","QRP 005 Discharge to community, risk-standardized (%)"),("qrp_005_dtc_rs_rate_fn",FN("QRP 005 rate")),
 ("qrp_005_dtc_comp_perf","QRP 005 discharge-to-community vs national (CMS determination)"),("qrp_005_dtc_volume","QRP 005 eligible stays"),
 ("qrp_006_mspb_score","QRP 006 Medicare spending per beneficiary ratio (1.0 = national)"),("qrp_006_mspb_score_fn",FN("QRP 006")),("qrp_006_mspb_numb","QRP 006 episodes"),
 ("qrp_039_hai_rs_rate","QRP 039 Healthcare-associated infections requiring hospitalization, risk-standardized (%)"),("qrp_039_hai_rs_rate_fn",FN("QRP 039 rate")),
 ("qrp_039_hai_comp_perf","QRP 039 HAI vs national (CMS determination)"),
 ("qrp_013_obs_rate","QRP 013 Falls with major injury (%)"),("qrp_013_obs_rate_fn",FN("QRP 013")),
 ("qrp_038_adj_rate","QRP 038 New/worsened pressure ulcers (%, adjusted)"),("qrp_038_adj_rate_fn",FN("QRP 038")),
 ("qrp_042_obs_rate","QRP 042 At/above expected self-care & mobility at discharge (%)"),("qrp_042_obs_rate_fn",FN("QRP 042")),
 ("qrp_024_obs_rate","QRP 024 Self-care goal met (%)"),("qrp_024_obs_rate_fn",FN("QRP 024")),
 ("qrp_025_obs_rate","QRP 025 Mobility goal met (%)"),("qrp_025_obs_rate_fn",FN("QRP 025")),
 ("qrp_043_obs_rate","QRP 043 Med list to next provider (%)"),("qrp_043_obs_rate_fn",FN("QRP 043")),
 ("qrp_044_obs_rate","QRP 044 Med list to patient/family (%)"),("qrp_044_obs_rate_fn",FN("QRP 044")),
 ("qrp_007_obs_rate","QRP 007 Drug regimen review with follow-up (%)"),("qrp_007_obs_rate_fn",FN("QRP 007")),
 ("qrp_040_obs_rate","QRP 040 Staff COVID vaccine up to date (%)"),("qrp_040_obs_rate_fn",FN("QRP 040")),
 ("qrp_041_obs_rate","QRP 041 Staff flu vaccine (%)"),("qrp_041_obs_rate_fn",FN("QRP 041")),
 ("qrp_045_obs_rate","QRP 045 Residents COVID vaccine up to date (%)"),("qrp_045_obs_rate_fn",FN("QRP 045")),
 ("vbp_in_program","In SNF VBP FY2026 (Y/N)"),("vbp_ranking","SNF VBP FY2026 ranking (1 = best of 13,900)"),("vbp_performance_score","SNF VBP performance score (0–100)"),
 ("vbp_incentive_multiplier","SNF VBP incentive payment multiplier (1.0 = break-even; <1.0 = net Part A cut; floor 0.980322)"),
 ("x_vbp_net_part_a_payment_change_pct","x_ DERIVED: VBP net Part A payment change (%) = (multiplier − 1) × 100"),
 ("vbp_fy24_rsrr","SNF VBP FY24 risk-std readmission rate (SNFRM; fraction 0–1)"),("vbp_fy24_rsrr_fn",FN("VBP readmission")),
 ("vbp_fy24_hai_rate","SNF VBP FY24 risk-std HAI rate (fraction 0–1)"),("vbp_fy24_hai_rate_fn",FN("VBP HAI")),
 ("vbp_fy24_turnover","SNF VBP FY24 total nursing turnover (fraction 0–1)"),("vbp_fy24_turnover_fn",FN("VBP turnover")),
 ("vbp_fy24_adj_nurse_hprd","SNF VBP FY24 adjusted nurse hrs per resident-day"),("vbp_fy24_adj_nurse_hprd_fn",FN("VBP nurse hrs")),
 ("data_vintage","Data vintage (CMS processing date)"),
]
cols=[c for c,_ in CUR]; heads=[h for _,h in CUR]
assert len(set(heads))==len(heads), "duplicate headers: " + str([h for h in heads if heads.count(h)>1])
assert not any(h.startswith(" ") for h in heads)
master = df(f"SELECT {', '.join(cols)} FROM facility_master ORDER BY state, provider_name")
assert len(master)==14690
master.columns=heads
master_all = df("SELECT * FROM facility_master ORDER BY state, provider_name")
master_all.to_csv(f"{ROOT}/output/Facility_Master_AllColumns.csv", index=False)
NCOL = len(heads); NALL = master_all.shape[1]

# ---------- measure catalog with direction (our reading; flagged) ----------
cat = df("SELECT * FROM measure_catalog ORDER BY source, measure_code")
def direction(desc):
    if desc is None: return ""
    d=desc.lower()
    if any(k in d for k in ["vaccine","vaccin","flu shot","medication list","medications were reviewed","return to home","achieve or exceed","at or above"]): return "higher is better"
    if any(k in d for k in ["readmission","rehospitalized","emergency department","hospitalizations","falls","pressure ulcer","infection","spending","antipsychotic","antianxiety","restrained","weight","catheter","urinary","depressive","increased","worsened","incontinence"]): return "lower is better"
    return "UNCLEAR - verify"
dcol="x_better_direction (our interpretation of description; needs Quality Expert sign-off)"
cat[dcol]=cat["description"].map(direction)
cat.loc[cat["measure_code"].str.contains("COMP_PERF|VOLUME|NUMB|NUMERATOR|DENOMINATOR|_2_5|_97_5|NUMBER$|OBS_READM", regex=True), dcol]="n/a (count/CI/text)"
cat["raw_file"]=cat["source"].map({"mds_qm":"djen-97ju_2026-09-30_NH_QualityMsr_MDS_Sep2026_2.csv","claims_qm":"ijh5-nb2v_2026-09-30_NH_QualityMsr_Claims_Sep2026_2.csv","snf_qrp":"fykj-qjee_2026-09-30_Skilled_Nursing_Facility_Quality_Reporting_Program_Provider_Data_Sep2026_2.csv"})

# ---------- supporting detail ----------
penalties = df("SELECT ccn, provider_name, state, penalty_date, penalty_type, fine_amount, payment_denial_start_date, payment_denial_length_in_days FROM penalties ORDER BY ccn, penalty_date")
cit_severe = df("""SELECT ccn, provider_name, state, survey_date, survey_type, inspection_cycle, deficiency_tag_number, deficiency_description, deficiency_category, scope_severity_code,
                   standard_deficiency, complaint_deficiency, infection_control_inspection_deficiency, deficiency_corrected, correction_date, citation_under_idr, citation_under_iidr
                   FROM health_citations WHERE scope_severity_code IN ('G','H','I','J','K','L') ORDER BY ccn, survey_date""")
ownership = df("""SELECT ccn, provider_name, state, role_played_by_owner_or_manager_in_facility AS role, owner_type, owner_name, ownership_percentage AS ownership_pct_numeric, ownership_percentage_raw, association_date, association_date_raw
                  FROM ownership WHERE role LIKE '%OWNERSHIP INTEREST%' OR role LIKE '%OPERATIONAL/MANAGERIAL CONTROL%' OR role LIKE '%PARTNERSHIP%' OR role='Ownership Data Not Available' ORDER BY ccn, role""")
_vraw = [r[0] for r in con.execute("SELECT column_name FROM (DESCRIBE vbp_facility) WHERE column_name LIKE '%_raw'").fetchall()]
vbp = df(f"""SELECT CASE WHEN p.ccn IS NOT NULL THEN 'Y' ELSE 'N' END AS in_nursing_home_universe, v.* EXCLUDE ({', '.join(_vraw)}) FROM vbp_facility v LEFT JOIN provider_info p USING (ccn) ORDER BY snf_vbp_program_ranking""")
state_avg = df("SELECT * FROM state_averages")
cut = df("SELECT * FROM inspection_cutpoints")
intervals = df("SELECT * FROM collection_intervals")
castlog = pd.read_csv(f"{ROOT}/db/cast_log.csv")
lineage = pd.read_csv(f"{ROOT}/db/column_map.csv")

# ---------- lineage for Facility_Master: every header → raw file + original header / formula ----------
SRC = {"provider_info":"4pq5-n9py_2026-09-30_NH_ProviderInfo_Sep2026_3.csv","survey_summary":"tbry-pc2d_2026-09-30_NH_SurveySummary_Sep2026_2.csv",
       "health_citations":"r5ix-sfxw_2026-09-30_NH_HealthCitations_Sep2026_2.csv","fire_citations":"ifjz-ge4w_2026-09-30_NH_FireSafetyCitations_Sep2026_2.csv",
       "survey_dates":"svdt-c123_2026-09-30_NH_SurveyDates_Sep2026_2.csv","penalties":"g6vv-u9sr_2026-09-30_NH_Penalties_Sep2026_2.csv",
       "ownership":"y2hd-n93e_2026-09-30_NH_Ownership_Sep2026_2.csv","mds":"djen-97ju_2026-09-30_NH_QualityMsr_MDS_Sep2026_2.csv",
       "claims":"ijh5-nb2v_2026-09-30_NH_QualityMsr_Claims_Sep2026_2.csv","qrp":"fykj-qjee_2026-09-30_Skilled_Nursing_Facility_Quality_Reporting_Program_Provider_Data_Sep2026_2.csv",
       "vbp":"284v-j9fz_2025-11-01_FY_2026_SNF_VBP_Facility_Performance.csv"}
pi_map = {r.column:r.original_header for r in lineage[lineage.table=="provider_info"].itertuples()}
qrp_codes = {r[0]:r[1] for r in con.execute("SELECT master_column, measure_code FROM measure_catalog WHERE source='snf_qrp'").fetchall()}
def lin(c):
    if c in pi_map: return (SRC["provider_info"], pi_map[c], "direct copy")
    if c=="data_vintage": return (SRC["provider_info"], "Processing Date", "direct copy (renamed)")
    if c.startswith("ss_c1_"): return (SRC["survey_summary"], {"ss_c1_health_survey_date":"Health Survey Date","ss_c1_health_deficiencies":"Total Number of Health Deficiencies","ss_c1_fire_deficiencies":"Total Number of Fire Safety Deficiencies"}[c], "direct copy of the Inspection Cycle = 1 row")
    m=re.match(r"qm_(\d{3})_(4q|adj|obs|exp|fn)$", c)
    if m:
        code, kind = m.groups()
        if code in ("521","522","551","552"):
            return (SRC["claims"], {"adj":"Adjusted Score","obs":"Observed Score","exp":"Expected Score","fn":"Footnote for Score"}[kind], f"row where Measure Code = {code}")
        return (SRC["mds"], {"4q":"Four Quarter Average Score","fn":"Footnote for Four Quarter Average Score"}[kind], f"row where Measure Code = {code}")
    if c.startswith("qrp_"):
        base = c[:-3] if c.endswith("_fn") else c
        return (SRC["qrp"], "Footnote" if c.endswith("_fn") else "Score", f"row where Measure Code = {qrp_codes[base]}" + (" ('Not Available' → blank)" if not c.endswith("_fn") else " ('-' → blank)"))
    if c.startswith("vbp_") or c.startswith("x_vbp"):
        vm = {"vbp_ranking":"SNF VBP Program Ranking","vbp_performance_score":"Performance Score","vbp_incentive_multiplier":"Incentive Payment Multiplier",
              "vbp_fy24_rsrr":"Performance Period: FY 2024 Risk-Standardized Readmission Rate","vbp_fy24_hai_rate":"Performance Period: FY 2024 Risk-Standardized Healthcare-Associated Infection Rate",
              "vbp_fy24_turnover":"Performance Period: FY 2024 Total Nursing Staff Turnover Rate","vbp_fy24_adj_nurse_hprd":"Performance Period: FY 2024 Adjusted Total Nursing Staff Hours per Resident Day",
              "vbp_in_program":"(presence of the CCN in the file)","x_vbp_net_part_a_payment_change_pct":"Incentive Payment Multiplier"}
        if c.endswith("_fn"): return (SRC["vbp"], "Footnote -- " + vm[c[:-3]], "'---' → blank; long CMS sentences shortened (see LINEAGE note)")
        note = "DERIVED: (multiplier − 1) × 100" if c.startswith("x_") else ("'---' → blank" if c.startswith("vbp_fy24") else "direct copy")
        return (SRC["vbp"], vm[c], note)
    agg = {
     "hc_n_total":("health_citations","count of rows"),"hc_n_severity_G_or_worse":("health_citations","count of rows with Scope Severity Code in G,H,I,J,K,L"),
     "hc_n_actual_harm_GHI":("health_citations","count with Scope Severity Code in G,H,I"),"hc_n_immediate_jeopardy_JKL":("health_citations","count with Scope Severity Code in J,K,L"),
     "hc_n_abuse_neglect":("health_citations","count with Deficiency Category = 'Freedom from Abuse, Neglect, and Exploitation Deficiencies'"),
     "hc_n_standard":("health_citations","count with Standard Deficiency = Y"),"hc_n_complaint":("health_citations","count with Complaint Deficiency = Y"),
     "hc_n_complaint_last_12m":("health_citations","count with Complaint Deficiency = Y and Survey Date ≥ 2025-09-01"),
     "hc_n_complaint_last_12m_G_or_worse":("health_citations","count with Complaint Deficiency = Y, Survey Date ≥ 2025-09-01, severity G–L"),
     "hc_n_infection_control":("health_citations","count with Infection Control Inspection Deficiency = Y"),
     "hc_n_last_12m":("health_citations","count with Survey Date ≥ 2025-09-01"),"hc_n_last_12m_G_or_worse":("health_citations","count with Survey Date ≥ 2025-09-01 and severity G–L"),
     "fc_n_total":("fire_citations","count of rows"),
     "sd_n_health_complaint_surveys_3cycles":("survey_dates","count with Type of Survey = 'Health Complaint' (CMS lists complaint inspections only if they produced a citation)"),
     "sd_n_health_complaint_surveys_last_12m":("survey_dates","count with Type of Survey = 'Health Complaint' and Survey Date ≥ 2025-09-01"),
     "sd_n_infection_control_surveys":("survey_dates","count with Type of Survey = 'Infection Control'"),
     "pen_n_fines":("penalties","count with Penalty Type = Fine"),"pen_total_fines_usd":("penalties","sum of Fine Amount where Penalty Type = Fine"),
     "pen_n_fines_last_12m":("penalties","count Fine with Penalty Date ≥ 2025-09-01"),"pen_fines_usd_last_12m":("penalties","sum Fine Amount with Penalty Date ≥ 2025-09-01"),
     "pen_n_payment_denials":("penalties","count with Penalty Type = Payment Denial"),"pen_total_denial_days":("penalties","sum of Payment Denial Length in Days"),
     "pen_last_penalty_date":("penalties","max Penalty Date"),
     "own_data_available":("ownership","N if the facility's rows say 'Ownership Data Not Available', else Y"),
     "own_direct_owners":("ownership","distinct Owner Name where Role contains 'DIRECT OWNERSHIP INTEREST' (not INDIRECT), joined with ' | '"),
     "own_indirect_owners":("ownership","distinct Owner Name where Role contains 'INDIRECT OWNERSHIP INTEREST'"),
     "own_partnership_owners":("ownership","distinct Owner Name where Role contains 'PARTNERSHIP INTEREST'"),
     "own_operational_control_orgs":("ownership","distinct Owner Name where Role = 'OPERATIONAL/MANAGERIAL CONTROL' and Owner Type = Organization"),
     "own_operational_control_individuals":("ownership","distinct Owner Name where Role = 'OPERATIONAL/MANAGERIAL CONTROL' and Owner Type = Individual"),
     "own_most_recent_association_date":("ownership","max Association Date ('since MM/DD/YYYY' parsed; 'NO DATE PROVIDED' → blank)"),
    }
    if c in agg: t,f = agg[c]; return (SRC[t], "(aggregate)", f"{f} — scripts/02_load.py")
    raise KeyError(c)
lineage_cur = pd.DataFrame([{"Facility_Master header":h, "db column":c, **dict(zip(["source_file","original_header","how_derived"], lin(c)))} for c,h in CUR])
assert lineage_cur["source_file"].notna().all()
lineage_cur.loc[len(lineage_cur)] = {"Facility_Master header":"(note)","db column":"vbp_*_fn","source_file":SRC["vbp"],"original_header":"Footnote -- …",
  "how_derived":"CMS text shortened: 'did not meet this measure's case minimum policy requirement…' → 'Case minimum not met - not reported'; 'did not have sufficient data to calculate a baseline period measure result' → 'Insufficient data for baseline period' (+ '/ improvement score' variant). Full text is in db/snf.duckdb vbp_facility.*_raw."}
recon = json.load(open(f"{ROOT}/db/reconciliation.json"))
recon_df = pd.DataFrame([{"Check":k,"Result":str(v)} for k,v in recon.items()])

readme = pd.DataFrame({"SNF Intelligence — CMS Care Compare, Sep 2026 refresh (workbook v2)":[
 "WHAT THIS IS: every CMS-certified nursing home (14,690 facilities) with ratings, staffing, quality measures, deficiencies, complaints, penalties, ownership, SNF QRP and SNF VBP results joined into one row per facility.",
 "SOURCE: data.cms.gov/provider-data/topics/nursing-homes. Nursing-home files: processing date 2026-09-01 (the 2026-09-30 release). SNF QRP provider file: the 2026-09-30 release (no processing-date column; measure periods on Collection_Intervals). SNF VBP FY2026 facility file dated 2025-11-01. The Jul-2026 QRP national/swing-bed files are not used in this workbook.",
 f"HOW TO USE: Facility_Master ({NCOL} columns, plain-English headers, formatted as an Excel Table named FacilityMaster) is the sheet to filter/pivot. The companion CSV Facility_Master_AllColumns.csv has all {NALL} database columns. Detail sheets hold row-level penalties, severe citations and ownership records.",
 "CSV WARNING: if you open Facility_Master_AllColumns.csv by double-clicking, Excel strips leading zeros from CCN and ZIP (025025 becomes 25025). Use Data > From Text/CSV and set those columns to Text, or just use this workbook.",
 "FOOTNOTES: every value CMS can suppress has a '— footnote' column, normally immediately to its right. Exceptions: one 'Reported staffing hours — footnote' covers all reported staffing-hour columns; case-mix-adjusted hours are explained by the Staffing stars footnote; QRP 'eligible stays' / 'vs national' blanks are explained by that measure's rate footnote a few columns to the left. A blank value with a footnote means CMS suppressed or did not compute it (e.g. footnote 9 = too few residents). Blank is NOT zero. Footnote meanings: docs/DATA_DICTIONARY_DIGEST.md section 7.",
 "COUNTS OF ZERO: for penalties, citations and complaint inspections, 0 means none appear in CMS's published window (3 inspection cycles / ~3 years), which is exactly what CMS publishes.",
 "COMPLAINTS: CMS publishes complaint-inspection deficiencies and the dates of complaint inspections that resulted in a citation. It does NOT publish the number of complaints received or substantiated, so those are not here.",
 "NOTHING IS ESTIMATED OR IMPUTED. Every column traces to a CMS column (LINEAGE_FacilityMaster) or to a counted/summed aggregate of CMS rows whose rule is written on that sheet. The one derived column is prefixed x_ and its formula is in its header.",
 "WINDOWS: 'last 12 months' = Survey/Penalty Date on or after 2025-09-01, which is CMS's own 'most recent 12 month period' (Collection_Intervals, C_H_CMPL). '3 cycles' / '3-yr' = everything CMS publishes in the file.",
 "SEVERITY: Scope/severity letters A–L. G–I = actual harm; J–L = immediate jeopardy. The CMS data dictionary does not define the full grid (its Abuse Icon definition does say severity 'G or greater' means harm); the full definition is in the CMS Five-Star Quality Rating System Technical Users' Guide, which this project has not yet independently fetched — treat the G–I / J–L labels as [industry context] until it is.",
 "SNF VBP (the 'withhold'): CMS withholds 2% of each SNF's Medicare Part A fee-for-service payments and returns 60% of the pool as incentive payments based on performance (cms.gov SNF VBP program page, checked 2026-10-26). The incentive payment multiplier is applied to the facility's Part A claims: 1.0 = break-even, below 1.0 = net cut; the floor is 0.980322 (a 1.97% net cut, i.e. essentially the whole 2% withhold). In FY2026, 10,843 of the 13,775 nursing homes in the program have a multiplier below 1.0. Dollar amounts are not yet published by CMS ('to be added at the end of FY 2026').",
 "UNITS: Provider Info turnover and all QM/QRP rates are percentages (0–100). SNF VBP rates are fractions (0–1) exactly as CMS publishes them; headers say which.",
 "RATING vs INSPECTION CYCLE: Provider Info's 'Rating Cycle 1' deficiency counts differ from Survey Summary 'Inspection Cycle 1' counts for 449 facilities (plus 38 unrated): 287 explained by CMS excluding citations under informal dispute, 162 not yet explained. This workbook uses Survey Summary / citation-file counts everywhere.",
 "VERIFICATION: an independent red-team agent (different model, did not build the workbook) compared every Facility_Master column against the raw CMS files for all 14,690 facilities (all matched) and 5 facilities against the live CMS provider-data API (all matched; same data vintage). Reports: docs/VALIDATION_REPORT_01.md and _02.md. The medicare.gov Care Compare web page could not be machine-read (JavaScript), so a human spot-check there is still recommended.",
 f"Built {datetime.date.today()} by scripts/02_load.py + scripts/03_export_xlsx.py from db/snf.duckdb. Every type conversion is logged on the CAST_LOG sheet with proof that no non-blank value was lost.",
]})
limits = pd.DataFrame({"Known limitations":[
 "1. No hospital→SNF discharge flow or referral-source information exists in any CMS Care Compare nursing-home file. Catchment / 'where do a hospital's patients go' analysis requires a different source (docs/EXTERNAL_DATA_SOURCES.md): no free public source has true flows; MedPAR (paid, research DUA) does; a modeled catchment can be built from free files.",
 "2. No cost or revenue fields. Spending signals: the SNF QRP Medicare Spending per Beneficiary ratio (relative to national, not dollars) and the SNF VBP incentive multiplier (a % change to Part A payments, not dollars).",
 "3. Quality measures are facility-level averages over a reporting window (Measure_Catalog / Collection_Intervals); they do not describe individual patients or condition-specific outcomes.",
 "4. MDS quality measures are self-reported by facilities via MDS assessments; claims-based measures (QM 521/522/551/552; SNF QRP 004/005/006/039; VBP readmission and HAI) come from Medicare claims; VBP turnover and staffing come from Payroll-Based Journal data.",
 "5. Reporting windows differ by measure and are older than the refresh date: SNF QRP 004/005/006 cover 10/2022–9/2024; SNF VBP FY2026 uses FY2024 performance; most MDS QMs 2025Q2–2026Q1 (flu-vaccine QMs 454/472: 2024Q3–2025Q2; short-stay 430/434 start 2025Q1); claims QMs CY2025. See Collection_Intervals.",
 "6. Provider Info 'Rating Cycle 1' deficiency counts differ from Survey Summary 'Inspection Cycle 1' counts for 449 facilities; 287 explained (citations under dispute), 162 unexplained — open item for the Quality Expert.",
 "7. Ownership text columns on Facility_Master are our concatenation of CMS ownership rows by role; 637 facilities have no ownership data (flag column). Use Ownership_Detail for exact records, including the raw percentage text ('NOT APPLICABLE' / 'NO PERCENTAGE PROVIDED').",
 "8. Health-inspection star ratings are relative within state (Inspection_Cutpoints) — a 3-star in Maryland is not the same deficiency score as a 3-star in Kentucky.",
 "9. Medicaid-only facilities (247) are not in the SNF QRP/VBP programs (QRP footnote 14), but CMS does publish claims-based QMs for many of them (167 of 247 have QM 551) because those measures use all Medicare claims for their residents.",
 "10. Complaint counts here are deficiencies from complaint inspections and complaint inspections that produced citations. The number of complaints filed or substantiated is not public in these files.",
 "11. The A–L severity grid labels (actual harm / immediate jeopardy) are industry-standard but not defined in the CMS data dictionary; they are pending verification against the CMS Five-Star Technical Users' Guide.",
 "12. Any 'x_' column is OUR derivation, not a CMS value: x_vbp_net_part_a_payment_change_pct on Facility_Master and the direction column on Measure_Catalog (which still needs Quality Expert sign-off).",
]})

print("writing...")
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    readme.to_excel(xw, "README", index=False)
    master.to_excel(xw, "Facility_Master", index=False)
    cat.to_excel(xw, "Measure_Catalog", index=False)
    penalties.to_excel(xw, "Penalties_Detail", index=False)
    cit_severe.to_excel(xw, "Citations_SeverityG-L", index=False)
    ownership.to_excel(xw, "Ownership_Detail", index=False)
    vbp.to_excel(xw, "VBP_FY2026", index=False)
    state_avg.to_excel(xw, "State_Averages", index=False)
    cut.to_excel(xw, "Inspection_Cutpoints", index=False)
    intervals.to_excel(xw, "Collection_Intervals", index=False)
    lineage_cur.to_excel(xw, "LINEAGE_FacilityMaster", index=False)
    lineage.to_excel(xw, "LINEAGE_AllColumns", index=False)
    recon_df.to_excel(xw, "RECONCILIATION", index=False)
    castlog.to_excel(xw, "CAST_LOG", index=False)
    limits.to_excel(xw, "LIMITATIONS", index=False)

print("formatting...")
wb=load_workbook(OUT)
wb._named_styles['Normal'].font = Font(name="Arial", size=10)
hdr_fill=PatternFill("solid", fgColor="1F3864"); hdr_font=Font(name="Arial", size=10, bold=True, color="FFFFFF")
for ws in wb.worksheets:
    ws.freeze_panes="B2" if ws.title=="Facility_Master" else "A2"
    for c in ws[1]:
        c.fill=hdr_fill; c.font=hdr_font; c.alignment=Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height=58 if ws.max_column>8 else 18
    if ws.max_row>1 and ws.title not in ("README","LIMITATIONS","Facility_Master"):
        ws.auto_filter.ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for i in range(1, ws.max_column+1):
        ws.column_dimensions[get_column_letter(i)].width = 110 if ws.title in ("README","LIMITATIONS") else (15 if ws.max_column>20 else 30)
    if ws.title in ("README","LIMITATIONS"):
        for r in ws.iter_rows(min_row=2):
            for c in r: c.alignment=Alignment(wrap_text=True, vertical="top")
ws=wb["Facility_Master"]
tbl=Table(displayName="FacilityMaster", ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
tbl.tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True); ws.add_table(tbl)
DATE_COLS = {"own_most_recent_association_date","date_first_approved_to_provide_medicare_and_medicaid_services","ss_c1_health_survey_date","pen_last_penalty_date","data_vintage"}
for j,(c,h) in enumerate(CUR, start=1):
    fmt = '$#,##0' if "$" in h else ('yyyy-mm-dd' if c in DATE_COLS else None)
    if fmt == 'yyyy-mm-dd': assert "%" not in h and str(master.dtypes.iloc[j-1]).startswith("datetime"), (c, h, master.dtypes.iloc[j-1])
    if fmt:
        for r in range(2, ws.max_row+1): ws.cell(r,j).number_format=fmt
assert all(c in DATE_COLS for c,h in CUR if str(master[h].dtype).startswith("datetime")), "a datetime column is missing from DATE_COLS"
# Summary sheet with live formulas
ss=wb.create_sheet("Summary", 1)
ss["A1"]="Live summary (formulas over Facility_Master)"; ss["A1"].font=Font(name="Arial", bold=True, size=12)
ss["A3"]="Overall star rating"; ss["B3"]="Facilities"; ss["C3"]="Avg QM 521 short-stay rehospitalization %"; ss["D3"]="Avg total fines $ (3-yr)"; ss["E3"]="Avg VBP multiplier"
L=lambda name: get_column_letter(cols.index(name)+1)
for i,star in enumerate([1,2,3,4,5], start=4):
    ss[f"A{i}"]=star
    ss[f"B{i}"]=f"=COUNTIF(Facility_Master!${L('overall_rating')}:${L('overall_rating')},A{i})"
    ss[f"C{i}"]=f"=IFERROR(AVERAGEIF(Facility_Master!${L('overall_rating')}:${L('overall_rating')},A{i},Facility_Master!${L('qm_521_adj')}:${L('qm_521_adj')}),\"\")"
    ss[f"D{i}"]=f"=IFERROR(AVERAGEIF(Facility_Master!${L('overall_rating')}:${L('overall_rating')},A{i},Facility_Master!${L('pen_total_fines_usd')}:${L('pen_total_fines_usd')}),\"\")"
    ss[f"E{i}"]=f"=IFERROR(AVERAGEIF(Facility_Master!${L('overall_rating')}:${L('overall_rating')},A{i},Facility_Master!${L('vbp_incentive_multiplier')}:${L('vbp_incentive_multiplier')}),\"\")"
    ss[f"C{i}"].number_format='0.00'; ss[f"D{i}"].number_format='$#,##0'; ss[f"E{i}"].number_format='0.0000'
ss["A9"]="Not rated"; ss["B9"]=f"=COUNTBLANK(Facility_Master!${L('overall_rating')}$2:${L('overall_rating')}$14691)"
ss["A10"]="Total"; ss["B10"]="=SUM(B4:B9)"; ss["B10"].font=Font(name="Arial", bold=True)
ss["A12"]="Note: simple facility averages (unweighted), matching how CMS computes State_Averages. Blank (suppressed) values are excluded by AVERAGEIF."
for c in ("A3","B3","C3","D3","E3"): ss[c].font=Font(name="Arial", bold=True)
for col,w in zip("ABCDE",(20,12,40,24,18)): ss.column_dimensions[col].width=w
wb.save(OUT)
print("saved", OUT, "| Facility_Master columns:", NCOL)
